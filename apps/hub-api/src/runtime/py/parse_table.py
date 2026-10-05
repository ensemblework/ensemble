"""Binary table readers, run by the person's own python3. Same behaviour as agent-runtime's parse_table.py."""

from __future__ import annotations

import io
from typing import Any

MAX_ROWS = 200_000
PICKLE = "Pickle files can run code when they are opened. Ensemble will not load them. Export the frame as Parquet or Feather instead."


def _frame(columns: list[str], rows: list[list[Any]], sheets: list[str], sheet: str) -> dict[str, Any]:
    typed = _infer(columns, rows)
    warnings: list[str] = []
    if len(rows) > MAX_ROWS:
        warnings.append(f"Kept the first {MAX_ROWS} rows.")
        rows = rows[:MAX_ROWS]
    return {"columns": typed, "rows": rows, "sheets": sheets, "sheet": sheet, "warnings": warnings}


def _infer(names: list[str], rows: list[list[Any]]) -> list[dict[str, str]]:
    columns = []
    for index, name in enumerate(names):
        seen = 0
        numeric = 0
        for row in rows[:200]:
            if index >= len(row):
                continue
            value = row[index]
            if value is None or value == "":
                continue
            seen += 1
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                numeric += 1
        kind = "number" if seen and numeric / seen >= 0.8 else "text"
        if kind == "text":
            unique = {str(row[index]) for row in rows[:200] if index < len(row) and row[index] not in (None, "")}
            if unique and len(unique) <= max(12, seen * 0.2):
                kind = "category"
        columns.append({"name": str(name)[:200] or f"column_{index + 1}", "type": kind})
    return columns


def _records(frame: Any) -> tuple[list[str], list[list[Any]]]:
    import pandas as pd

    if not isinstance(frame, pd.DataFrame):
        raise ValueError("That file did not contain a table.")
    names = [str(column) for column in frame.columns]
    rows: list[list[Any]] = []
    for record in frame.itertuples(index=False, name=None):
        row: list[Any] = []
        for value in record:
            if value is None or (isinstance(value, float) and value != value):
                row.append(None)
            elif hasattr(value, "isoformat"):
                row.append(value.isoformat())
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                row.append(float(value) if isinstance(value, float) else int(value))
            else:
                row.append(str(value))
        rows.append(row)
    return names, rows


def parse_table(filename: str, data: bytes, sheet: str = "") -> dict[str, Any]:
    name = filename.lower()
    if name.endswith(".pkl") or name.endswith(".pickle"):
        raise ValueError(PICKLE)
    if name.endswith(".parquet"):
        import pyarrow.parquet as pq

        frame = pq.read_table(io.BytesIO(data)).to_pandas()
        columns, rows = _records(frame)
        return _frame(columns, rows, [], "")
    if name.endswith(".feather") or name.endswith(".arrow"):
        import pyarrow.feather as feather

        frame = feather.read_feather(io.BytesIO(data))
        columns, rows = _records(frame)
        return _frame(columns, rows, [], "")
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_vba=False)
        sheets = list(book.sheetnames)
        chosen = sheet if sheet in sheets else (sheets[0] if sheets else "")
        ws = book[chosen] if chosen else None
        matrix = []
        if ws is not None:
            for record in ws.iter_rows(values_only=True):
                if any(cell is not None and str(cell).strip() for cell in record):
                    matrix.append(list(record))
        book.close()
        return _from_matrix(matrix, sheets, chosen)
    if name.endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        sheets = book.sheet_names()
        chosen = sheet if sheet in sheets else (sheets[0] if sheets else "")
        ws = book.sheet_by_name(chosen)
        matrix = [[ws.cell_value(r, c) for c in range(ws.ncols)] for r in range(ws.nrows)]
        return _from_matrix(matrix, sheets, chosen)
    if name.endswith(".ods"):
        import pandas as pd

        book = pd.ExcelFile(io.BytesIO(data), engine="odf")
        sheets = list(book.sheet_names)
        chosen = sheet if sheet in sheets else (sheets[0] if sheets else "")
        frame = book.parse(chosen)
        columns, rows = _records(frame)
        return _frame(columns, rows, sheets, chosen)
    if name.endswith(".numbers"):
        import tempfile
        from numbers_parser import Document

        with tempfile.NamedTemporaryFile(suffix=".numbers") as handle:
            handle.write(data)
            handle.flush()
            document = Document(handle.name)
            sheets = [item.name for item in document.sheets]
            chosen_sheet = document.sheets[0]
            if sheet:
                for item in document.sheets:
                    if item.name == sheet:
                        chosen_sheet = item
                        break
            table = chosen_sheet.tables[0]
            matrix = table.rows(values_only=True)
            return _from_matrix([list(row) for row in matrix], sheets, chosen_sheet.name)
    raise ValueError("That file type is not supported by the plot runtime.")


def _from_matrix(matrix: list[list[Any]], sheets: list[str], sheet: str) -> dict[str, Any]:
    if not matrix:
        return _frame([], [], sheets, sheet)
    header = [str(cell).strip() if cell is not None else "" for cell in matrix[0]]
    body = matrix[1:] if any(header) else matrix
    if not any(header):
        header = [f"column_{index + 1}" for index in range(len(matrix[0]))]
    width = max(len(header), max((len(row) for row in body), default=0))
    names = [(header[i] if i < len(header) and header[i] else f"column_{i + 1}") for i in range(width)]
    rows: list[list[Any]] = []
    for raw in body:
        row: list[Any] = []
        for index in range(width):
            value = raw[index] if index < len(raw) else None
            if value is None or value == "":
                row.append(None)
            elif isinstance(value, bool):
                row.append(int(value))
            elif isinstance(value, (int, float)):
                row.append(int(value) if isinstance(value, float) and value.is_integer() else value)
            elif hasattr(value, "isoformat"):
                row.append(value.isoformat())
            else:
                text = str(value).strip()
                row.append(text or None)
        if any(cell is not None for cell in row):
            rows.append(row)
    return _frame(names, rows, sheets, sheet)
