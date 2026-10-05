"""Parser, sandbox, and publication PDF tests for Plots."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

from ensemble_agent.plots.parse_table import PICKLE, parse_table
from ensemble_agent.plots.sandbox import run_plot

FIXTURES = Path(__file__).parent / "fixtures"
FIXTURES.mkdir(exist_ok=True)


def _sales_rows():
    return [
        ["month", "revenue", "cost"],
        ["Jan", 12000, 8000],
        ["Feb", 15000, 9000],
        ["Mar", 18000, 11000],
    ]


def test_pickle_is_refused():
    import pickle

    blob = pickle.dumps({"a": 1})
    (FIXTURES / "frame.pkl").write_bytes(blob)
    with pytest.raises(ValueError, match="Pickle"):
        parse_table("frame.pkl", blob)
    assert "Parquet" in PICKLE


def test_xlsx_and_xlsm_roundtrip():
    import openpyxl

    for name in ("sales.xlsx", "sales.xlsm"):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Revenue"
        for row in _sales_rows():
            sheet.append(row)
        extra = book.create_sheet("Empty")
        extra.append(["note"])
        buffer = io.BytesIO()
        book.save(buffer)
        path = FIXTURES / name
        path.write_bytes(buffer.getvalue())
        parsed = parse_table(name, path.read_bytes(), "Revenue")
        assert parsed["sheet"] == "Revenue"
        assert "Empty" in parsed["sheets"]
        assert parsed["columns"][0]["name"] == "month"
        assert parsed["rows"][0][1] == 12000


def test_xls_roundtrip():
    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("Revenue")
    for r, row in enumerate(_sales_rows()):
        for c, value in enumerate(row):
            sheet.write(r, c, value)
    buffer = io.BytesIO()
    book.save(buffer)
    (FIXTURES / "sales.xls").write_bytes(buffer.getvalue())
    parsed = parse_table("sales.xls", buffer.getvalue())
    assert parsed["rows"][1][0] == "Feb"
    assert parsed["rows"][1][1] == 15000


def test_ods_roundtrip():
    import pandas as pd

    frame = pd.DataFrame(
        [{"month": "Jan", "revenue": 12000}, {"month": "Feb", "revenue": 15000}],
    )
    path = FIXTURES / "sales.ods"
    frame.to_excel(path, engine="odf", index=False)
    parsed = parse_table("sales.ods", path.read_bytes())
    assert parsed["columns"][0]["name"] == "month"
    assert parsed["rows"][0][1] == 12000


def test_numbers_roundtrip():
    from numbers_parser import Document

    document = Document()
    table = document.sheets[0].tables[0]
    rows = _sales_rows()
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.write(r, c, value)
    path = FIXTURES / "sales.numbers"
    document.save(str(path))
    parsed = parse_table("sales.numbers", path.read_bytes())
    assert parsed["rows"][0][0] == "Jan"
    assert parsed["rows"][0][1] == 12000


def test_parquet_and_feather():
    import pyarrow as pa
    import pyarrow.feather as feather
    import pyarrow.parquet as pq

    table = pa.table({"month": ["Jan", "Feb"], "revenue": [12000, 15000]})
    parquet = FIXTURES / "sales.parquet"
    feather_path = FIXTURES / "sales.feather"
    pq.write_table(table, parquet)
    feather.write_feather(table, feather_path)
    for name in (parquet.name, feather_path.name):
        parsed = parse_table(name, (FIXTURES / name).read_bytes())
        assert parsed["columns"][1]["type"] == "number"
        assert parsed["rows"][1][1] == 15000


def _code(body: str) -> str:
    return "from ensemble_plots import load, save, apply_publication\nimport matplotlib.pyplot as plt\n" + body


def test_matplotlib_pdf_embeds_type42():
    code = _code(
        """
apply_publication("icml")
df = load("sales")
fig, ax = plt.subplots(figsize=(3.25, 2.4))
ax.plot(df["month"], df["revenue"], color="#0072B2")
ax.axhline(10000, color="#D55E00", linestyle="dashed", label="expected")
ax.set_xlabel("month")
ax.set_ylabel("revenue")
fig.tight_layout()
save(fig)
"""
    )
    result = run_plot(
        code,
        {"datasets": []} if False else [{"name": "sales", "columns": ["month", "revenue"], "rows": [["Jan", 12000], ["Feb", 15000], ["Mar", 18000]]}],
    )
    assert not result.get("error"), result.get("stderr")
    assert result["pdf"]
    import base64

    pdf = base64.b64decode(result["pdf"])
    assert pdf.startswith(b"%PDF")
    assert b"FontFile2" in pdf
    assert result["svg"] and "<svg" in result["svg"]
    assert result["png"]
    assert result["eps"]


def test_pickle_import_stays_blocked_and_missing_data_is_reported():
    blocked = run_plot("import pickle\n", [])
    assert blocked.get("error")
    assert "not allowed" in (blocked.get("stderr") or "")
    missing = run_plot("from ensemble_plots import load\nload('nope')\n", [{"name": "sales", "columns": ["a"], "rows": [[1]]}])
    assert missing.get("error")
    assert "sales" in (missing.get("stderr") or "")


def test_os_confinement_and_git_token(monkeypatch):
    """Desktop and hosted plots share one runner. Guards block files, network, and processes. macOS also applies Seatbelt."""
    import os

    from ensemble_agent.plots.sandbox import reset_plot_worker

    monkeypatch.setenv("ENSEMBLE_DESKTOP", "1")
    token = "ghp_plot_escape_token_value"
    os.environ["ENSEMBLE_GIT_TOKEN"] = token
    reset_plot_worker()
    try:
        probed = run_plot(
            "import os\n"
            "print('TOKEN=' + os.environ.get('ENSEMBLE_GIT_TOKEN', ''))\n"
            "print(open('/etc/passwd').read(20))\n",
            [],
        )
    finally:
        os.environ.pop("ENSEMBLE_GIT_TOKEN", None)
        reset_plot_worker()
    blob = (probed.get("stdout") or "") + (probed.get("stderr") or "")
    assert token not in blob
    assert probed.get("error")
    assert "root:" not in (probed.get("stdout") or "")


@pytest.mark.skipif(__import__("sys").platform == "darwin", reason="macOS always uses the spawn choke point")
def test_hosted_runner_blocks_network_files_processes_and_secrets(monkeypatch):
    """Hosted Linux has no OS sandbox, so main's in-process guards (#33) stay on."""
    import os

    monkeypatch.delenv("ENSEMBLE_DESKTOP", raising=False)
    monkeypatch.setenv("ENSEMBLE_GIT_TOKEN", "ghp_hosted_plot_token_value")
    monkeypatch.setenv("OPENAI_API_KEY", "sk-hosted-plot-key-value")
    from ensemble_agent.plots.sandbox import reset_plot_worker

    reset_plot_worker()
    network = run_plot("import socket\nsocket.create_connection(('example.com', 80))\n", [])
    assert network.get("error")
    files = run_plot("print(open('/etc/passwd').read())\n", [])
    assert files.get("error")
    assert "root:" not in (files.get("stdout") or "")
    blocked = run_plot("import pickle\n", [])
    assert blocked.get("error")
    spawn = run_plot("import os\nos.system('echo unsandboxed')\n", [])
    assert spawn.get("error")
    assert "unsandboxed" not in (spawn.get("stdout") or "")
    secrets = run_plot("import os\nprint(sorted(os.environ))\nprint(os.environ.get('ENSEMBLE_GIT_TOKEN', ''), os.environ.get('OPENAI_API_KEY', ''))\n", [])
    blob = (secrets.get("stdout") or "") + (secrets.get("stderr") or "")
    assert "ghp_hosted_plot_token_value" not in blob
    assert "sk-hosted-plot-key-value" not in blob
    assert "OPENAI_API_KEY" not in blob
    assert os.environ.get("ENSEMBLE_GIT_TOKEN") == "ghp_hosted_plot_token_value"


def test_plot_mode_follows_the_desktop_flag_and_the_os_sandbox(monkeypatch):
    import sys

    from ensemble_agent.plots import sandbox

    monkeypatch.setenv("ENSEMBLE_DESKTOP", "1")
    assert sandbox.plot_mode() == "sandbox"
    monkeypatch.delenv("ENSEMBLE_DESKTOP")
    assert sandbox.plot_mode() == ("sandbox" if sys.platform == "darwin" else "hosted")


def test_hosted_uses_the_warm_worker_and_macos_uses_the_spawn_choke_point(monkeypatch):
    from ensemble_agent.plots import sandbox

    seen = []

    def submit(job, timeout, on_started=None):
        seen.append((job, timeout))
        if on_started is not None:
            on_started()
        return {"stdout": "", "stderr": "", "png": "ok"}

    monkeypatch.setattr(sandbox._client, "submit", submit)
    monkeypatch.setattr(sandbox, "_warm", True)
    monkeypatch.setenv("ENSEMBLE_GIT_TOKEN", "ghp_mode_test_token_value")
    monkeypatch.setattr(sandbox, "plot_mode", lambda: "hosted")
    assert not sandbox.run_plot("print(1)\n", [], fmt="png", dpi=150).get("error")
    job, timeout = seen[-1]
    assert job["format"] == "png" and job["dpi"] == 150
    assert timeout == sandbox.WARM_TIMEOUT_S
    assert "ENSEMBLE_GIT_TOKEN" not in sandbox._worker_env()

    monkeypatch.setattr(sandbox, "plot_mode", lambda: "sandbox")
    assert not sandbox.run_plot("print(1)\n", [], fmt="png", dpi=150).get("error")
    assert seen[-1][1] == sandbox.TIMEOUT_S

    hosted = sandbox.worker_launch_args("hosted")
    assert hosted[-1].endswith("worker.py")
    assert not any(part.endswith("cli.ts") for part in hosted)
    seatbelt = sandbox.seatbelt_worker_argv("python3", "/tmp/worker.py", "/tmp/plots", ["/usr"], ["/cache"])
    assert "--worker" in seatbelt
    assert any(part.endswith("workspace/sandbox/cli.ts") for part in seatbelt)
    assert "sys.addaudithook" in sandbox._worker_file().read_text()


def _figure(token: str, y: int) -> str:
    return (
        "from ensemble_plots import save\n"
        "import matplotlib.pyplot as plt\n"
        f"print({token!r}, flush=True)\n"
        "fig, ax = plt.subplots()\n"
        f"ax.plot([1, 2], [1, {y}])\n"
        "save(fig)\n"
    )


def test_seatbelt_read_paths_come_from_the_plot_interpreter():
    import os
    import sys

    from ensemble_agent.plots.sandbox import _python_read_paths

    paths = _python_read_paths(sys.executable)
    prefix = os.path.realpath(sys.prefix)
    assert paths
    assert any(path == prefix or path.startswith(prefix + os.sep) for path in paths)


def test_bundled_plot_worker_matches_the_runtime_source():
    source = Path(__file__).resolve().parents[1] / "ensemble_agent" / "plots" / "worker.py"
    bundled = Path(__file__).resolve().parents[2] / "hub-api" / "src" / "runtime" / "py" / "plot_worker.py"
    assert source.read_bytes() == bundled.read_bytes()


def test_parent_stays_on_agg_without_threads():
    import json
    import subprocess
    import sys

    from ensemble_agent.plots.sandbox import _worker_env, _worker_file

    proc = subprocess.Popen(
        [sys.executable, str(_worker_file())],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        env=_worker_env(),
        text=True,
    )
    assert proc.stdout is not None
    line = proc.stdout.readline()
    proc.kill()
    proc.wait(timeout=5)
    ready = json.loads(line)
    assert ready["ready"] is True
    assert str(ready["backend"]).lower() == "agg"
    assert ready["threads"] == 1
    if __import__("sys").platform != "win32":
        assert ready["fork"] is True


def test_warm_export_is_under_a_second():
    import time

    from ensemble_agent.plots.sandbox import reset_plot_worker

    reset_plot_worker()
    code = _figure("warm", 2)
    assert not run_plot(code, [], fmt="png").get("error")
    started = time.perf_counter()
    again = run_plot(code, [], fmt="png")
    elapsed = time.perf_counter() - started
    assert not again.get("error"), again.get("stderr")
    assert again.get("png")
    assert elapsed < 1.0, elapsed


def test_overlapping_exports_keep_their_own_pngs():
    import threading

    codes = [_figure("TOKEN-A", 2), _figure("TOKEN-B", 9)]
    results: list[dict | None] = [None, None]

    def run(index: int) -> None:
        results[index] = run_plot(codes[index], [], fmt="png")

    threads = [threading.Thread(target=run, args=(index,)) for index in range(2)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert results[0] and results[1]
    assert "TOKEN-A" in (results[0].get("stdout") or "")
    assert "TOKEN-B" in (results[1].get("stdout") or "")
    assert "TOKEN-A" not in (results[1].get("stdout") or "")
    assert results[0].get("png") and results[0]["png"] != results[1].get("png")
    assert not results[0].get("error") and not results[1].get("error")


def test_export_past_the_cap_waits(monkeypatch):
    import threading
    import time

    from ensemble_agent.plots.sandbox import reset_plot_worker

    monkeypatch.setenv("ENSEMBLE_PLOT_CONCURRENCY", "2")
    reset_plot_worker()
    results: list[dict | None] = [None, None, None]

    def run(index: int) -> None:
        code = f"import time\nprint('START {index}', time.time(), flush=True)\ntime.sleep(0.8)\nprint('END {index}', flush=True)\n"
        results[index] = run_plot(code, [], fmt="png")

    threads = [threading.Thread(target=run, args=(index,)) for index in range(3)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert all(result and not result.get("error") for result in results)
    starts = []
    for result in results:
        assert result is not None
        for line in (result.get("stdout") or "").splitlines():
            if line.startswith("START"):
                starts.append(float(line.split()[-1]))
    assert len(starts) == 3
    ordered = sorted(starts)
    assert ordered[1] - ordered[0] < 0.5
    assert ordered[2] - ordered[0] > 0.5


def test_queue_time_is_not_the_timeout_and_one_timeout_is_alone(monkeypatch):
    import threading

    from ensemble_agent.plots import sandbox
    from ensemble_agent.plots.sandbox import reset_plot_worker

    monkeypatch.setattr(sandbox, "WARM_TIMEOUT_S", 3)
    monkeypatch.setattr(sandbox, "TIMEOUT_S", 3)
    monkeypatch.setenv("ENSEMBLE_PLOT_CONCURRENCY", "1")
    reset_plot_worker()
    monkeypatch.setattr(sandbox, "_warm", True)
    queued: list[dict | None] = [None, None, None]

    def run(index: int) -> None:
        queued[index] = run_plot("import time\ntime.sleep(2)\nprint('OK', flush=True)\n", [])

    threads = [threading.Thread(target=run, args=(index,)) for index in range(3)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert all(result and not result.get("error") and "OK" in (result.get("stdout") or "") for result in queued)

    monkeypatch.setenv("ENSEMBLE_PLOT_CONCURRENCY", "2")
    reset_plot_worker()
    monkeypatch.setattr(sandbox, "_warm", True)
    warm = run_plot("print('WARM', flush=True)\n", [])
    assert "WARM" in (warm.get("stdout") or "")
    proc = sandbox._client.proc
    assert proc is not None and proc.poll() is None
    pid = proc.pid
    pair: list[dict | None] = [None, None]

    def slow() -> None:
        pair[0] = run_plot("import time\ntime.sleep(30)\nprint('SLOW', flush=True)\n", [])

    def quick() -> None:
        pair[1] = run_plot("import time\ntime.sleep(2)\nprint('QUICK', flush=True)\n", [])

    slow_thread = threading.Thread(target=slow)
    quick_thread = threading.Thread(target=quick)
    slow_thread.start()
    quick_thread.start()
    quick_thread.join(timeout=20)
    slow_thread.join(timeout=20)
    assert pair[1] and not pair[1].get("error") and "QUICK" in (pair[1].get("stdout") or "")
    assert pair[0] and "longer than" in (pair[0].get("error") or "")
    assert "SLOW" not in (pair[0].get("stdout") or "")
    assert sandbox._client.proc is not None and sandbox._client.proc.pid == pid and sandbox._client.proc.poll() is None
    still = run_plot("print('STILL', flush=True)\n", [])
    assert "STILL" in (still.get("stdout") or "") and not still.get("error")
    assert sandbox._client.proc is not None and sandbox._client.proc.pid == pid


def test_killing_the_worker_stops_its_child():
    import os
    import signal
    import sys
    import threading
    import time

    from ensemble_agent.plots import sandbox
    from ensemble_agent.plots.sandbox import reset_plot_worker

    if sys.platform == "win32":
        return
    reset_plot_worker()
    box: dict = {}

    def run() -> None:
        box["result"] = run_plot("import time\ntime.sleep(60)\nprint('DONE', flush=True)\n", [])

    thread = threading.Thread(target=run)
    thread.start()
    child = None
    worker_pid = None
    for _ in range(100):
        proc = sandbox._client.proc
        if proc is not None and proc.poll() is None:
            worker_pid = proc.pid
            kids = _child_pids(worker_pid)
            if kids:
                child = kids[0]
                break
        time.sleep(0.05)
    assert child is not None and worker_pid is not None
    os.kill(worker_pid, signal.SIGKILL)
    gone = False
    for _ in range(50):
        if not _running(child):
            gone = True
            break
        time.sleep(0.1)
    thread.join(timeout=10)
    reset_plot_worker()
    assert gone, child


def _child_pids(parent: int) -> list[int]:
    import os

    found = []
    for name in os.listdir("/proc"):
        if not name.isdigit():
            continue
        try:
            text = open(f"/proc/{name}/status").read()
        except OSError:
            continue
        for line in text.splitlines():
            if line.startswith("PPid:") and int(line.split()[1]) == parent:
                found.append(int(name))
    return found


def _running(pid: int) -> bool:
    try:
        text = open(f"/proc/{pid}/status").read()
    except OSError:
        return False
    for line in text.splitlines():
        if line.startswith("State:"):
            return line.split()[1] not in {"Z", "X"}
    return False


def test_stale_plot_directories_are_swept(tmp_path):
    import os
    import subprocess
    import sys
    import time

    from ensemble_agent.plots.worker import _sweep_stale_plots

    stale = tmp_path / "ensemble-plot-stale"
    fresh = tmp_path / "ensemble-plot-fresh"
    held = tmp_path / "ensemble-plot-held"
    for path in (stale, fresh, held):
        path.mkdir()
    old = time.time() - 600
    os.utime(stale, (old, old))
    os.utime(held, (old, old))
    proc = subprocess.Popen([sys.executable, "-c", "import os, sys, time; os.chdir(sys.argv[1]); time.sleep(30)", str(held)])
    try:
        for _ in range(50):
            try:
                cwd = os.path.realpath(f"/proc/{proc.pid}/cwd")
            except OSError:
                cwd = ""
            if cwd == os.path.realpath(held):
                break
            time.sleep(0.05)
        _sweep_stale_plots(tmp_path, max_age=180)
    finally:
        proc.kill()
        proc.wait(timeout=5)
    assert not stale.exists()
    assert fresh.exists()
    assert held.exists()


def test_infinite_loop_hits_the_cpu_limit():
    import sys
    import time

    if sys.platform == "win32":
        return
    started = time.perf_counter()
    result = run_plot("while True:\n    pass\n", [])
    elapsed = time.perf_counter() - started
    assert elapsed < 15, elapsed
    assert "CPU" in (result.get("error") or "")


def test_large_allocation_is_refused():
    import sys

    if sys.platform == "win32":
        return
    result = run_plot("x = bytearray(3 * 1024 * 1024 * 1024)\nprint('allocated', len(x))\n", [])
    blob = (result.get("error") or "") + (result.get("stderr") or "")
    assert "MemoryError" in blob
    assert "allocated" not in (result.get("stdout") or "")


def test_write_over_32mb_fails():
    import sys

    if sys.platform == "win32":
        return
    result = run_plot("open('big.bin','wb').write(b'x' * (33 * 1024 * 1024))\nprint('WROTE')\n", [])
    assert "file size" in (result.get("error") or "").lower()
    assert "WROTE" not in (result.get("stdout") or "")


def test_killing_the_worker_between_exports_lets_the_next_one_run():
    from ensemble_agent.plots import sandbox

    first = run_plot("print('BEFORE-KILL', flush=True)\n", [])
    assert "BEFORE-KILL" in (first.get("stdout") or "")
    sandbox._client.shutdown()
    second = run_plot("print('AFTER-KILL', flush=True)\n", [])
    assert "AFTER-KILL" in (second.get("stdout") or "")
    assert not second.get("error")


def test_fork_crash_falls_back_to_a_fresh_process(monkeypatch):
    import sys

    from ensemble_agent.plots.sandbox import reset_plot_worker

    if sys.platform == "win32":
        return
    monkeypatch.setenv("ENSEMBLE_PLOT_TEST_CRASH", "fork")
    reset_plot_worker()
    try:
        result = run_plot("print('AFTER-CRASH', flush=True)\n", [])
    finally:
        monkeypatch.delenv("ENSEMBLE_PLOT_TEST_CRASH", raising=False)
        reset_plot_worker()
    assert "AFTER-CRASH" in (result.get("stdout") or "")
    assert not result.get("error")


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (None, 2),
        ("", 2),
        ("  ", 2),
        ("abc", 2),
        ("1.5", 2),
        ("1e2", 2),
        ("0", 1),
        ("-3", 1),
        ("1", 1),
        ("2", 2),
        ("32", 32),
        ("33", 32),
        ("100", 32),
    ],
)
def test_plot_concurrency_parses_like_the_desktop(monkeypatch, raw, expected):
    from ensemble_agent.config import _plot_concurrency
    from ensemble_agent.plots.sandbox import plot_concurrency
    from ensemble_agent.plots.worker import concurrency

    if raw is None:
        monkeypatch.delenv("ENSEMBLE_PLOT_CONCURRENCY", raising=False)
    else:
        monkeypatch.setenv("ENSEMBLE_PLOT_CONCURRENCY", raw)
    assert plot_concurrency() == expected
    assert concurrency() == expected
    assert _plot_concurrency() == expected


def test_queue_start_delay_uses_the_limits_ahead_and_the_cap():
    from ensemble_agent.plots.sandbox import queue_start_delay

    assert queue_start_delay([], 1) == 0
    assert queue_start_delay([4.5], 1) == 4.5
    assert queue_start_delay([4.5, 4.5], 1) == 9
    assert queue_start_delay([4.5, 4.5], 2) == 4.5
    assert queue_start_delay([4.5, 4.5, 2], 2) == 4.5


def test_queue_wait_follows_the_first_run_ahead_not_the_waiters_own_limit(monkeypatch):
    import threading
    import time

    from ensemble_agent.plots import sandbox
    from ensemble_agent.plots.sandbox import PLOT_EXPORTS_BUSY, reset_plot_worker

    # Production, cap 1: timeout, then a 1.5s export, a 35s export, and quick X.
    # X's own warm limit is 20s, so the old wait was 25s. Scaled to 1/10, slack included.
    monkeypatch.setattr(sandbox, "WARM_TIMEOUT_S", 2)
    monkeypatch.setattr(sandbox, "TIMEOUT_S", 2)
    monkeypatch.setattr(sandbox, "COLD_TIMEOUT_S", 4.5)
    monkeypatch.setattr(sandbox, "QUEUE_SLACK_S", 0.5)
    monkeypatch.setenv("ENSEMBLE_PLOT_CONCURRENCY", "1")
    reset_plot_worker()
    monkeypatch.setattr(sandbox, "_warm", True)
    forced = run_plot("import time\ntime.sleep(30)\nprint('FORCE', flush=True)\n", [])
    assert "longer than" in (forced.get("error") or ""), forced
    assert sandbox._warm is False
    box: dict[str, dict] = {}

    def run(name: str, code: str) -> None:
        box[name] = run_plot(code, [])

    short = threading.Thread(target=run, args=("short", "import time\ntime.sleep(0.4)\nprint('SHORT', flush=True)\n"))
    long = threading.Thread(target=run, args=("long", "import time\ntime.sleep(3.5)\nprint('LONG', flush=True)\n"))
    short.start()
    long.start()
    for _ in range(100):
        with sandbox._client.lock:
            if len(sandbox._client.pending) >= 2:
                break
        time.sleep(0.02)
    with sandbox._client.lock:
        assert len(sandbox._client.pending) >= 2
        assert all(item.timeout == 4.5 for item in sandbox._client.pending.values())
    short.join(timeout=10)
    assert "SHORT" in (box["short"].get("stdout") or ""), box.get("short")
    assert not box["short"].get("error")
    assert sandbox._warm is True
    started = time.perf_counter()
    result = run_plot("print('X', flush=True)\n", [])
    elapsed = time.perf_counter() - started
    long.join(timeout=15)
    error = str(result.get("error") or "")
    assert "could not draw" not in error
    assert "did not start" not in error
    assert error in {"", PLOT_EXPORTS_BUSY}
    assert "X" in (result.get("stdout") or ""), result
    assert not result.get("error")
    # The old waiter returned at 2.5s. The 3.5s export was still inside its 4.5s limit.
    assert elapsed > 3, elapsed
    assert "LONG" in (box["long"].get("stdout") or ""), box.get("long")
    assert not box["long"].get("error")


def test_queue_give_up_is_busy_not_a_figure_error(tmp_path, monkeypatch):
    from ensemble_agent.plots import sandbox
    from ensemble_agent.plots.sandbox import PLOT_EXPORTS_BUSY, reset_plot_worker

    stub = tmp_path / "python-never-starts"
    stub.write_text(
        "#!/usr/bin/env python3\n"
        "import json, sys\n"
        "sys.stdout.write(json.dumps({\"ready\": True, \"backend\": \"Agg\", \"threads\": 1, \"fork\": False}) + \"\\n\")\n"
        "sys.stdout.flush()\n"
        "for line in sys.stdin:\n"
        "    pass\n"
    )
    stub.chmod(0o755)
    monkeypatch.setenv("ENSEMBLE_PYTHON", str(stub))
    monkeypatch.setattr(sandbox, "QUEUE_SLACK_S", 0.4)
    reset_plot_worker()
    try:
        result = run_plot("print('X', flush=True)\n", [])
    finally:
        reset_plot_worker()
    assert result.get("error") == PLOT_EXPORTS_BUSY
    assert "could not draw" not in (result.get("error") or "")
    assert "did not start" not in (result.get("error") or "")
    assert result.get("error") != "The plot runtime is not available."


def test_queue_wait_cap_returns_busy_before_the_proxy_cut(monkeypatch):
    """Concurrency 1 and enough work ahead that the real wait exceeds the cap.

    The cap is configurable so this does not wait the production 110s. The
    waiter still gets the busy error at or before that short cap.
    """
    import threading
    import time

    from ensemble_agent.plots import sandbox
    from ensemble_agent.plots.sandbox import PLOT_EXPORTS_BUSY, QUEUE_WAIT_CAP_S, reset_plot_worker

    assert QUEUE_WAIT_CAP_S <= 110
    cap = 0.35
    monkeypatch.setattr(sandbox, "QUEUE_WAIT_CAP_S", cap)
    monkeypatch.setattr(sandbox, "QUEUE_SLACK_S", 0.05)
    monkeypatch.setattr(sandbox, "WARM_TIMEOUT_S", 3)
    monkeypatch.setattr(sandbox, "TIMEOUT_S", 3)
    monkeypatch.setattr(sandbox, "COLD_TIMEOUT_S", 3)
    monkeypatch.setenv("ENSEMBLE_PLOT_CONCURRENCY", "1")
    reset_plot_worker()
    monkeypatch.setattr(sandbox, "_warm", True)
    box: dict[str, dict] = {}

    def run(name: str, code: str) -> None:
        box[name] = run_plot(code, [])

    ahead = [
        threading.Thread(target=run, args=("first", "import time\ntime.sleep(2)\nprint('A', flush=True)\n")),
        threading.Thread(target=run, args=("second", "import time\ntime.sleep(2)\nprint('B', flush=True)\n")),
    ]
    for thread in ahead:
        thread.start()
    for _ in range(100):
        with sandbox._client.lock:
            if len(sandbox._client.pending) >= 2:
                break
        time.sleep(0.02)
    with sandbox._client.lock:
        assert len(sandbox._client.pending) >= 2
    started = time.perf_counter()
    result = run_plot("print('X', flush=True)\n", [])
    elapsed = time.perf_counter() - started
    for thread in ahead:
        thread.join(timeout=15)
    try:
        assert result.get("error") == PLOT_EXPORTS_BUSY, result
        assert "X" not in (result.get("stdout") or "")
        # Two 3s limits would be 6s. The waiter must stop at the cap.
        assert elapsed <= cap + 0.5, elapsed
        assert elapsed < 2, elapsed
    finally:
        reset_plot_worker()


def test_alarm_backstop_survives_alarm_zero(monkeypatch):
    import time

    from ensemble_agent.plots import sandbox
    from ensemble_agent.plots.sandbox import reset_plot_worker

    monkeypatch.setattr(sandbox, "WARM_TIMEOUT_S", 3)
    monkeypatch.setattr(sandbox, "TIMEOUT_S", 3)
    reset_plot_worker()
    monkeypatch.setattr(sandbox, "_warm", True)
    started = time.perf_counter()
    result = run_plot("import signal, time\nsignal.alarm(0)\ntime.sleep(30)\nprint('STILL', flush=True)\n", [])
    elapsed = time.perf_counter() - started
    assert "longer than" in (result.get("error") or ""), result
    assert "STILL" not in (result.get("stdout") or "")
    assert elapsed < 15, elapsed


def _pid_alive(pid: int) -> bool:
    import os

    try:
        os.kill(pid, 0)
    except OSError:
        return False
    try:
        text = open(f"/proc/{pid}/status").read()
    except OSError:
        return False
    return "\nState:\tZ" not in text and not text.startswith("State:\tZ")


def test_parent_sigkill_stops_alarm_mask_and_reloaded_guard(monkeypatch):
    """In-process guards can be disarmed. The parent's group SIGKILL still fires."""
    import os
    import time

    from ensemble_agent.plots import sandbox
    from ensemble_agent.plots.sandbox import reset_plot_worker

    if os.name != "posix":
        return
    monkeypatch.setattr(sandbox, "WARM_TIMEOUT_S", 2)
    monkeypatch.setattr(sandbox, "TIMEOUT_S", 2)
    monkeypatch.setattr(sandbox, "COLD_TIMEOUT_S", 2)
    reset_plot_worker()
    try:
        warm = run_plot("print('WARM', flush=True)\n", [])
        assert "WARM" in (warm.get("stdout") or ""), warm
        monkeypatch.setattr(sandbox, "_warm", True)
        scripts = {
            "alarm": (
                "import os, signal, time\n"
                "signal.alarm(0)\n"
                "print('PID', os.getpid(), flush=True)\n"
                "child = os.fork()\n"
                "if child == 0:\n"
                "    print('CHILD', os.getpid(), flush=True)\n"
                "    time.sleep(120)\n"
                "    os._exit(0)\n"
                "print('FORKED', child, flush=True)\n"
                "time.sleep(120)\n"
                "print('STILL', flush=True)\n"
            ),
            "mask": (
                "import os, signal, time\n"
                "signal.pthread_sigmask(signal.SIG_BLOCK, {signal.SIGALRM, signal.SIGTERM, signal.SIGINT})\n"
                "print('PID', os.getpid(), flush=True)\n"
                "time.sleep(120)\n"
                "print('STILL', flush=True)\n"
            ),
            "reload": (
                "import importlib, os, signal, time\n"
                "importlib.reload(signal)\n"
                "signal.alarm(0)\n"
                "print('PID', os.getpid(), flush=True)\n"
                "time.sleep(120)\n"
                "print('STILL', flush=True)\n"
            ),
        }
        for name, code in scripts.items():
            started = time.perf_counter()
            result = run_plot(code, [])
            elapsed = time.perf_counter() - started
            stdout = result.get("stdout") or ""
            assert "longer than" in (result.get("error") or ""), (name, result)
            assert "STILL" not in stdout, (name, result)
            assert elapsed < 3.5, (name, elapsed, result)
            pids = [int(line.split()[1]) for line in stdout.splitlines() if line.startswith(("PID ", "CHILD ", "FORKED "))]
            assert pids, (name, stdout)
            for _ in range(20):
                if not any(_pid_alive(pid) for pid in pids):
                    break
                time.sleep(0.1)
            alive = [pid for pid in pids if _pid_alive(pid)]
            assert alive == [], (name, alive)
    finally:
        reset_plot_worker()


def test_python_override_without_matplotlib_is_unavailable(tmp_path, monkeypatch):
    from ensemble_agent.plots.sandbox import reset_plot_worker

    stub = tmp_path / "python-no-mpl"
    stub.write_text(
        "#!/usr/bin/env python3\nimport sys\nsys.stdout.write('{\"ready\": false, \"missing\": \"matplotlib\"}\\n')\nsys.stdout.flush()\n"
    )
    stub.chmod(0o755)
    monkeypatch.setenv("ENSEMBLE_PYTHON", str(stub))
    reset_plot_worker()
    try:
        result = run_plot("print(1)\n", [])
    finally:
        reset_plot_worker()
    assert result.get("error") == "The plot runtime is not available."
